"""Read an ExxonMobil crude assay workbook (Summary + Yield Graph sheets)."""
import openpyxl

LIGHT_ENDS = {  # assay label -> DWSIM compound name
    "methane + ethane": "Ethane",
    "propane": "Propane",
    "isobutane": "Isobutane",
    "n-butane": "N-butane",
    "isopentane": "Isopentane",
    "n-pentane": "N-pentane",
}


def _num(x):
    return isinstance(x, (int, float))


def read_assay(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Summary (C)"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    name = next(r[5] for r in rows if len(r) > 5 and r[4] == "Crude:")
    whole = {}
    light = {}
    for r in rows:
        if len(r) > 11 and isinstance(r[7], str) and r[7].strip() in LIGHT_ENDS and _num(r[11]):
            light[LIGHT_ENDS[r[7].strip()]] = float(r[11])
        if len(r) > 16 and isinstance(r[12], str) and _num(r[16]):
            whole[r[12].strip()] = float(r[16])

    def row(label):
        return next(r for r in rows if len(r) > 1 and isinstance(r[1], str) and r[1].strip().startswith(label))

    start, end = row("Start (°C)"), row("End (°C)")
    props = {
        "wt": row("Yield (% wt)"),
        "vol": row("Yield (% vol)"),
        "vabp": row("Volume Average B.P."),
        "d15": row("Density @ 15°C"),
        "uopk": row("UOPK"),
        "mw": row("Molecular Weight"),
        "sulfur": row("Total Sulfur (% wt)"),
    }
    cuts = []
    for j in range(3, len(start)):  # col 2 is whole crude
        if start[j] is None or end[j] is None or str(start[j]).strip() == "":
            continue
        cut = {"start": start[j], "end": end[j]}
        for k, r in props.items():
            v = r[j] if j < len(r) else None
            cut[k] = float(v) if _num(v) else None
        cuts.append(cut)

    ws = wb["Yield Graph (C)"]
    tbp = []  # (T degC, cum wt %, cum vol %)
    col = None
    for r in ws.iter_rows(values_only=True):
        if col is None:
            if "Boiling Point" in r:
                col = r.index("Boiling Point")
            continue
        if len(r) > col + 2 and _num(r[col]) and _num(r[col + 1]):
            tbp.append((float(r[col]), float(r[col + 1]), float(r[col + 2])))
    return {"name": name, "whole": whole, "light_ends_wt": light, "cuts": cuts, "tbp": tbp}


if __name__ == "__main__":
    import sys
    a = read_assay(sys.argv[1])
    print(a["name"], {k: round(v, 4) for k, v in a["whole"].items() if k in ("Density @ 15°C (g/cc)", "API Gravity", "Total Sulfur (% wt)")})
    print("light ends wt%:", {k: round(v, 3) for k, v in a["light_ends_wt"].items()},
          "sum", round(sum(a["light_ends_wt"].values()), 3))
    print(f"{'cut':>12} {'wt%':>7} {'vol%':>7} {'VABP':>7} {'d15':>7} {'UOPK':>6} {'MW':>6} {'S%':>6}")
    for c in a["cuts"]:
        f = lambda v, p=2: "-" if v is None else f"{v:.{p}f}"
        print(f"{str(c['start'])+'-'+str(c['end']):>12} {f(c['wt']):>7} {f(c['vol']):>7} {f(c['vabp'],0):>7} "
              f"{f(c['d15'],4):>7} {f(c['uopk']):>6} {f(c['mw'],0):>6} {f(c['sulfur']):>6}")
    print("TBP points:", len(a["tbp"]), a["tbp"][0], a["tbp"][-1])
